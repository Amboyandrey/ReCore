"""Uploading a file into a conversation or a workflow, extracting whatever text it holds, and
reading it back.

Extraction is synchronous, on the request path, and covers what people actually attach: plain
text/markdown/JSON/CSV decoded directly, real parsing for the modern (XML-based) Office formats —
PDF, DOCX, PPTX, XLSX — and images, normalized (oriented, downscaled, HEIC transcoded to JPEG) and
handed to the model directly rather than OCR'd, since a vision-capable model reads layout,
diagrams and handwriting far better than OCR ever will. Anything else is stored but marked
unsupported: legacy binary Office formats (.doc/.xls/.ppt) need a much heavier tool (LibreOffice
headless, typically) than a pure-Python library, and audio/video need transcription, not text
extraction at all — genuinely different features, not a missing case of this one. "Attached" and
"extracted" are tracked separately, so adding either later is a new branch here, not a schema
change.
"""

import asyncio
import uuid
from io import BytesIO
from pathlib import Path

import pillow_heif
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.errors import AttachmentNotFound, AttachmentTooLarge
from app.models import Attachment, ExtractStatus

settings = get_settings()

# Lets Pillow's own Image.open() decode HEIC/HEIF transparently — the default format for photos
# straight off an iPhone, and by far the most likely image a user actually attaches.
pillow_heif.register_heif_opener()  # type: ignore[attr-defined]  # not in pillow_heif's __all__

_TEXT_MIME_PREFIXES = ("text/",)
_TEXT_MIME_TYPES = {"application/json", "application/xml", "application/x-yaml"}
_PDF_MIME = "application/pdf"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The intersection every provider adapter's vision input actually accepts (after normalization,
# HEIC/HEIF becomes JPEG — see _normalize_image — so they're accepted on the way in, not out).
_IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
_HEIC_MIMES = {"image/heic", "image/heif"}

# Anthropic documents no quality benefit above this on the long edge; the other providers cap
# similarly. Downscaling here means a 12MB phone photo reaches the model at all, rather than
# being rejected by the size cap below before it ever gets a chance to shrink.
MAX_IMAGE_DIMENSION = 1568

# Caps how much of one file's text reaches the model — a single huge document otherwise crowds
# out the rest of the conversation's context budget entirely.
MAX_EXTRACTED_CHARS = 100_000


def _is_text_like(mime: str) -> bool:
    """Whether this is read as plain UTF-8 text, no parsing library involved."""
    return mime.startswith(_TEXT_MIME_PREFIXES) or mime in _TEXT_MIME_TYPES


def _is_image(mime: str) -> bool:
    """Whether this is a format _normalize_image knows how to decode — PNG/JPEG/GIF/WebP
    directly, HEIC/HEIF via pillow-heif's opener registered above."""
    return mime in _IMAGE_MIMES or mime in _HEIC_MIMES


def _cap(text: str, max_chars: int | None) -> str:
    """Truncate extracted text at `max_chars`, with a visible marker that it happened. `None`
    means uncapped — used by knowledge indexing (see extract_text below), which chunks the full
    text itself rather than needing it pre-truncated to a context-budget-sized limit."""
    if max_chars is None or len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[...truncated]"


def _extract_plain_text(
    data: bytes, *, max_chars: int | None
) -> tuple[str | None, ExtractStatus, str | None]:
    """Text-like files are already text — just decode them."""
    try:
        return _cap(data.decode("utf-8"), max_chars), ExtractStatus.DONE, None
    except UnicodeDecodeError as exc:
        return None, ExtractStatus.FAILED, str(exc)


def _extract_pdf_text(data: bytes, *, max_chars: int | None) -> tuple[str | None, ExtractStatus, str | None]:
    """Concatenate every page's extracted text, page breaks marked with a blank line."""
    try:
        reader = PdfReader(BytesIO(data))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # noqa: BLE001 — pypdf raises several distinct errors for bad PDFs
        return None, ExtractStatus.FAILED, str(exc)
    if not text.strip():
        return None, ExtractStatus.FAILED, "No extractable text found (the PDF may be scanned images)."
    return _cap(text, max_chars), ExtractStatus.DONE, None


def _extract_docx_text(data: bytes, *, max_chars: int | None) -> tuple[str | None, ExtractStatus, str | None]:
    """Join every paragraph's text — tables and headers/footers aren't walked, just the body."""
    try:
        document = DocxDocument(BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
    except Exception as exc:  # noqa: BLE001 — python-docx raises several distinct errors too
        return None, ExtractStatus.FAILED, str(exc)
    if not text.strip():
        return None, ExtractStatus.FAILED, "No extractable text found."
    return _cap(text, max_chars), ExtractStatus.DONE, None


def _extract_pptx_text(data: bytes, *, max_chars: int | None) -> tuple[str | None, ExtractStatus, str | None]:
    """Join every slide's shape text, one slide per block — speaker notes aren't included."""
    try:
        presentation = Presentation(BytesIO(data))
        slides = []
        for slide in presentation.slides:
            lines = [
                shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
            ]
            slides.append("\n".join(line for line in lines if line))
        text = "\n\n".join(slide for slide in slides if slide)
    except Exception as exc:  # noqa: BLE001 — python-pptx raises several distinct errors too
        return None, ExtractStatus.FAILED, str(exc)
    if not text.strip():
        return None, ExtractStatus.FAILED, "No extractable text found."
    return _cap(text, max_chars), ExtractStatus.DONE, None


def _extract_xlsx_text(data: bytes, *, max_chars: int | None) -> tuple[str | None, ExtractStatus, str | None]:
    """Render each sheet as tab-separated rows — formula cells read as their cached value, if
    the file was saved with one; a formula never actually recalculated has nothing to show."""
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = [
                "\t".join("" if cell is None else str(cell) for cell in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sheets.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    except Exception as exc:  # noqa: BLE001 — openpyxl raises several distinct errors too
        return None, ExtractStatus.FAILED, str(exc)
    text = "\n\n".join(sheets)
    if not text.strip():
        return None, ExtractStatus.FAILED, "No extractable text found."
    return _cap(text, max_chars), ExtractStatus.DONE, None


def extract_text(
    data: bytes, mime: str, *, max_chars: int | None = MAX_EXTRACTED_CHARS
) -> tuple[str | None, ExtractStatus, str | None]:
    """Dispatch to the right extractor for this mime type, or mark it unsupported. Images never
    reach this — save_attachment routes them to _normalize_image before extraction would run.

    `max_chars` defaults to the same per-attachment budget chat context needs (see
    MAX_EXTRACTED_CHARS's own docstring); knowledge indexing (app/workers/index_connector.py)
    passes `None` since it chunks the full text itself rather than needing it pre-truncated.
    """
    if _is_text_like(mime):
        return _extract_plain_text(data, max_chars=max_chars)
    if mime == _PDF_MIME:
        return _extract_pdf_text(data, max_chars=max_chars)
    if mime == _DOCX_MIME:
        return _extract_docx_text(data, max_chars=max_chars)
    if mime == _PPTX_MIME:
        return _extract_pptx_text(data, max_chars=max_chars)
    if mime == _XLSX_MIME:
        return _extract_xlsx_text(data, max_chars=max_chars)
    return None, ExtractStatus.UNSUPPORTED, None


def _normalize_image(data: bytes, mime: str) -> tuple[bytes, str, ExtractStatus, str | None]:
    """Correct EXIF orientation, downscale to MAX_IMAGE_DIMENSION, and transcode HEIC/HEIF to
    JPEG — the one format every provider's vision input reliably accepts. Runs in a thread (see
    the caller): Pillow's decode/resize is CPU-bound and this process serves every other request
    too.

    On any decode failure (corrupt file, a decompression-bomb-sized image, ...) this falls back
    to storing the upload exactly as received, marked failed with a reason — the same "never
    reject the whole upload" contract every other extractor in this file follows.
    """
    try:
        opened = Image.open(BytesIO(data))
        image = ImageOps.exif_transpose(opened) or opened
        image.thumbnail((MAX_IMAGE_DIMENSION, MAX_IMAGE_DIMENSION))
        buffer = BytesIO()
        if mime == "image/png":
            image.save(buffer, format="PNG")
            return buffer.getvalue(), "image/png", ExtractStatus.PASSTHROUGH, None
        image.convert("RGB").save(buffer, format="JPEG", quality=85)
        return buffer.getvalue(), "image/jpeg", ExtractStatus.PASSTHROUGH, None
    except Exception as exc:  # noqa: BLE001 — Pillow/pillow-heif raise several distinct errors
        return data, mime, ExtractStatus.FAILED, str(exc)


async def save_attachment(
    db: AsyncSession,
    *,
    workspace_id: uuid.UUID,
    conversation_id: uuid.UUID | None = None,
    workflow_id: uuid.UUID | None = None,
    uploaded_by: uuid.UUID,
    filename: str,
    mime: str,
    data: bytes,
) -> Attachment:
    """Write an uploaded file to disk under the workspace, extract its text (or, for an image,
    normalize it for the model), and record both. Exactly one of `conversation_id` /
    `workflow_id` owns the file — the database's CHECK constraint enforces the same."""
    assert (conversation_id is None) != (workflow_id is None)  # mirrors attachments_one_parent
    extracted_text: str | None
    if _is_image(mime):
        data, mime, extract_status, extract_error = await asyncio.to_thread(
            _normalize_image, data, mime
        )
        extracted_text = None
    else:
        extracted_text, extract_status, extract_error = extract_text(data, mime)

    # Checked after normalization, not before: the whole point of downscaling is that a 10-12MB
    # phone photo reaches the model at all, rather than being rejected here before it ever
    # shrinks. Non-image files are untouched above, so this behaves exactly as it did before.
    if len(data) > settings.max_attachment_size_bytes:
        raise AttachmentTooLarge()

    attachment_id = uuid.uuid4()
    storage_key = f"{workspace_id}/{attachment_id}"
    path = Path(settings.storage_dir) / storage_key
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)

    attachment = Attachment(
        id=attachment_id,
        workspace_id=workspace_id,
        conversation_id=conversation_id,
        workflow_id=workflow_id,
        uploaded_by=uploaded_by,
        original_filename=filename,
        mime=mime,
        size=len(data),
        storage_key=storage_key,
        extracted_text=extracted_text,
        extract_status=extract_status,
        extract_error=extract_error,
    )
    db.add(attachment)
    await db.flush()
    return attachment


def read_attachment_bytes(attachment: Attachment) -> bytes:
    """Read an attachment's stored bytes back off disk — used to hand an image attachment to a
    provider adapter as an ImagePart. Synchronous, matching save_attachment's own direct
    `path.write_bytes()` call above: these are local files under a bind-mounted volume, not a
    network store, so there's no request-serving benefit to a thread hop here."""
    return (Path(settings.storage_dir) / attachment.storage_key).read_bytes()


async def list_attachments(db: AsyncSession, *, conversation_id: uuid.UUID) -> list[Attachment]:
    """List every attachment uploaded into a conversation, oldest first."""
    stmt = (
        select(Attachment)
        .where(Attachment.conversation_id == conversation_id)
        .order_by(Attachment.created_at)
    )
    return list((await db.scalars(stmt)).all())


async def get_attachment(
    db: AsyncSession, *, conversation_id: uuid.UUID, attachment_id: uuid.UUID
) -> Attachment:
    """Load one attachment by id, scoped to its conversation, or raise if it isn't there."""
    attachment = await db.scalar(
        select(Attachment).where(
            Attachment.id == attachment_id, Attachment.conversation_id == conversation_id
        )
    )
    if attachment is None:
        raise AttachmentNotFound()
    return attachment


async def attach_to_message(
    db: AsyncSession,
    *,
    conversation_id: uuid.UUID,
    attachment_ids: list[uuid.UUID],
    message_id: uuid.UUID,
) -> list[Attachment]:
    """Link already-uploaded attachments to the message that was just sent alongside them."""
    attachments = []
    for attachment_id in attachment_ids:
        attachment = await get_attachment(
            db, conversation_id=conversation_id, attachment_id=attachment_id
        )
        attachment.message_id = message_id
        attachments.append(attachment)
    await db.flush()
    return attachments


async def attach_to_run(
    db: AsyncSession,
    *,
    workflow_id: uuid.UUID,
    attachment_ids: list[uuid.UUID],
    run_id: uuid.UUID,
) -> list[Attachment]:
    """Link files uploaded into a workflow to the run that's starting with them — each must belong
    to this workflow and not already be spoken for by an earlier run."""
    attachments = []
    for attachment_id in attachment_ids:
        attachment = await db.scalar(
            select(Attachment).where(
                Attachment.id == attachment_id,
                Attachment.workflow_id == workflow_id,
                Attachment.workflow_run_id.is_(None),
            )
        )
        if attachment is None:
            raise AttachmentNotFound()
        attachment.workflow_run_id = run_id
        attachments.append(attachment)
    await db.flush()
    return attachments


async def list_run_attachments(db: AsyncSession, *, run_id: uuid.UUID) -> list[Attachment]:
    """Every file a workflow run was started with, oldest first."""
    stmt = (
        select(Attachment)
        .where(Attachment.workflow_run_id == run_id)
        .order_by(Attachment.created_at)
    )
    return list((await db.scalars(stmt)).all())


async def attachments_by_run_id(
    db: AsyncSession, run_ids: list[uuid.UUID]
) -> dict[uuid.UUID, list[Attachment]]:
    """The files behind each of several runs, grouped by run id — one query for a whole run list."""
    if not run_ids:
        return {}
    stmt = (
        select(Attachment)
        .where(Attachment.workflow_run_id.in_(run_ids))
        .order_by(Attachment.created_at)
    )
    grouped: dict[uuid.UUID, list[Attachment]] = {}
    for attachment in (await db.scalars(stmt)).all():
        assert attachment.workflow_run_id is not None  # the query above filtered on exactly that
        grouped.setdefault(attachment.workflow_run_id, []).append(attachment)
    return grouped
